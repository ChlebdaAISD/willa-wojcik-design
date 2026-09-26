#!/usr/bin/env python3
"""Kopia AIDS_company_data/umowy/_convert_to_docx.py dostosowana do umów Willa Wójcik
(numery ustępów widoczne, role podpisów Administrator/Procesor, podział strony, bez znaku middot).
Convert markdown contract files to elegantly formatted .docx for printing."""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


GRAY = RGBColor(0x88, 0x88, 0x88)
SIG_ROLE_RE = re.compile(r'^\*\*(Wykonawca|Zamawiający|Administrator|Procesor):\*\*\s*$')
PAGEBREAK = '<!-- pagebreak -->'
DARK = RGBColor(0x22, 0x22, 0x22)
BORDER_GRAY = '888888'
HEADER_BG = 'F2F2F2'


def _set_border(elem, edge, sz='4', color='000000', val='single'):
    border = OxmlElement(f'w:{edge}')
    border.set(qn('w:val'), val)
    border.set(qn('w:sz'), sz)
    border.set(qn('w:color'), color)
    elem.append(border)


def add_bottom_border_to_paragraph(paragraph, color=BORDER_GRAY, size='6'):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    _set_border(pBdr, 'bottom', sz=size, color=color)
    pPr.append(pBdr)


def shade_cell(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color_hex)
    shd.set(qn('w:val'), 'clear')
    tc_pr.append(shd)


def remove_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'nil')
        tblBorders.append(b)
    tblPr.append(tblBorders)


def add_page_number_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run1 = p.add_run()
    run1.font.size = Pt(9)
    run1.font.color.rgb = GRAY

    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    run1._r.append(fldChar1)
    instrText = OxmlElement('w:instrText')
    instrText.text = 'PAGE'
    run1._r.append(instrText)
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run1._r.append(fldChar2)

    p.add_run(' / ').font.size = Pt(9)

    run2 = p.add_run()
    run2.font.size = Pt(9)
    run2.font.color.rgb = GRAY
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'begin')
    run2._r.append(fldChar3)
    instrText2 = OxmlElement('w:instrText')
    instrText2.text = 'NUMPAGES'
    run2._r.append(instrText2)
    fldChar4 = OxmlElement('w:fldChar')
    fldChar4.set(qn('w:fldCharType'), 'end')
    run2._r.append(fldChar4)

    for run in p.runs:
        run.font.color.rgb = GRAY
        run.font.size = Pt(9)


def style_document(doc):
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10)
    style.paragraph_format.line_spacing = 1.2
    style.paragraph_format.space_after = Pt(3)

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    add_page_number_footer(doc)


def parse_inline(text):
    parts = []
    pattern = re.compile(r'\*\*(.+?)\*\*|`([^`]+)`')
    last = 0
    for m in pattern.finditer(text):
        if m.start() > last:
            parts.append((text[last:m.start()], 'plain'))
        if m.group(1) is not None:
            parts.append((m.group(1), 'bold'))
        else:
            parts.append((m.group(2), 'code'))
        last = m.end()
    if last < len(text):
        parts.append((text[last:], 'plain'))
    return parts


def render_runs(paragraph, text):
    for chunk, kind in parse_inline(text):
        run = paragraph.add_run(chunk)
        if kind == 'bold':
            run.bold = True
        elif kind == 'code':
            run.italic = True
            run.font.color.rgb = GRAY


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = DARK
    add_bottom_border_to_paragraph(p, color='444444', size='8')


def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = DARK
    add_bottom_border_to_paragraph(p, color=BORDER_GRAY, size='4')


def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = DARK


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    render_runs(p, text)
    return p


def add_numbered(doc, text, indent_level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5 + indent_level * 0.8)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(4)
    render_runs(p, text)
    return p


def add_bullet(doc, text, indent_level=0):
    p = doc.add_paragraph()
    base = 0.5 + indent_level * 0.6
    p.paragraph_format.left_indent = Cm(base)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(3)
    if indent_level == 0:
        prefix_run = p.add_run('— ')
        prefix_run.bold = True
        prefix_run.font.color.rgb = DARK
    else:
        prefix_run = p.add_run('– ')
        prefix_run.font.color.rgb = GRAY
    render_runs(p, text)
    return p


def add_checkbox(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(3)
    box_run = p.add_run('☐  ')
    box_run.bold = True
    render_runs(p, text)
    return p


def add_table_from_md(doc, lines):
    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip('|').split('|')]
        rows.append(cells)
    if len(rows) < 2:
        return
    rows.pop(1)
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = 'Table Grid'
    table.autofit = True

    for r, row_data in enumerate(rows):
        for c, cell_text in enumerate(row_data):
            if c >= cols:
                continue
            cell = table.cell(r, c)
            cell.text = ''
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for chunk, kind in parse_inline(cell_text):
                run = p.add_run(chunk)
                run.font.size = Pt(10)
                if r == 0 or kind == 'bold':
                    run.bold = True
                if kind == 'code':
                    run.italic = True
                    run.font.color.rgb = GRAY
            if r == 0:
                shade_cell(cell, HEADER_BG)
    doc.add_paragraph()


def parse_parties_block(lines, start_idx):
    """Parse the '## Strony umowy' block — returns (wykonawca_lines, zamawiajacy_lines, end_idx)."""
    wykonawca = []
    zamawiajacy = []
    current = None
    i = start_idx
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith('---'):
            return wykonawca, zamawiajacy, i
        if line.startswith('## ') or line.startswith('# '):
            return wykonawca, zamawiajacy, i
        stripped = line.lstrip()
        if stripped.startswith('**Wykonawca'):
            current = wykonawca
            i += 1
            continue
        if stripped.startswith('**Zamawiający'):
            current = zamawiajacy
            i += 1
            continue
        if current is not None and stripped.startswith('- '):
            current.append(stripped[2:].strip())
        i += 1
    return wykonawca, zamawiajacy, i


def add_parties_table(doc, wykonawca, zamawiajacy):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)

    table = doc.add_table(rows=1, cols=2)
    table.autofit = False

    hdr = table.rows[0]
    hdr.cells[0].text = ''
    hdr.cells[1].text = ''
    for idx, label in enumerate(['WYKONAWCA', 'ZAMAWIAJĄCY']):
        cell = hdr.cells[idx]
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(4)
        run = para.add_run(label)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = GRAY

    body_row = table.add_row()
    for idx, lines in enumerate([wykonawca, zamawiajacy]):
        cell = body_row.cells[idx]
        cell.text = ''
        first = True
        for line in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.paragraph_format.space_after = Pt(2)
            for chunk, kind in parse_inline(line):
                run = p.add_run(chunk)
                run.font.size = Pt(10)
                if kind == 'bold':
                    run.bold = True
                elif kind == 'code':
                    run.italic = True
                    run.font.color.rgb = GRAY

    remove_table_borders(table)
    doc.add_paragraph()


def add_signatures_table(doc, signatures):
    """signatures = list of dicts with keys 'role', 'name', 'date_label'."""
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False

    for idx, sig in enumerate(signatures[:2]):
        role_cell = table.rows[0].cells[idx]
        role_cell.text = ''
        p1 = role_cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(24)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run('_' * 32)
        r1.font.size = Pt(10)
        r1.font.color.rgb = GRAY

        p2 = role_cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(sig['role'])
        r2.bold = True
        r2.font.size = Pt(10)

        p3 = role_cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(0)
        r3 = p3.add_run(sig['name'])
        r3.font.size = Pt(9)
        r3.font.color.rgb = GRAY

        date_cell = table.rows[1].cells[idx]
        date_cell.text = ''
        pd = date_cell.paragraphs[0]
        pd.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pd.paragraph_format.space_before = Pt(8)
        pd.paragraph_format.space_after = Pt(0)
        rd = pd.add_run('Data: ')
        rd.font.size = Pt(9)
        rd.font.color.rgb = GRAY
        rd2 = pd.add_run(sig['date_label'])
        rd2.italic = True
        rd2.font.size = Pt(9)
        rd2.font.color.rgb = GRAY

    remove_table_borders(table)


def detect_signature_block(lines, start_idx):
    """Detect end-of-doc signature block. Returns (signatures, end_idx) or (None, start_idx)."""
    sigs = []
    i = start_idx
    current = None
    saw_signature = False
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.lstrip()

        role_m = SIG_ROLE_RE.match(stripped)
        if saw_signature and (stripped.startswith('#') or stripped.startswith('---') or stripped == PAGEBREAK):
            break
        if role_m:
            saw_signature = True
            role = role_m.group(1)
            current = {'role': role, 'name': '', 'date_label': '[DATA]'}
            sigs.append(current)
            i += 1
            continue

        if not saw_signature:
            return None, start_idx

        if current is not None:
            if 'Imię i Nazwisko:' in stripped or 'Nazwa:' in stripped:
                m = re.search(r'(?:Imię i Nazwisko|Nazwa):\s*(.+)', stripped)
                if m:
                    name_text = m.group(1).strip()
                    name_text = re.sub(r'`\[([^`]+)`?\]?', r'[\1]', name_text)
                    name_text = name_text.strip('`')
                    current['name'] = name_text
            elif 'Data:' in stripped:
                m = re.search(r'Data:\s*(.+)', stripped)
                if m:
                    d = m.group(1).strip().strip('`')
                    current['date_label'] = d
            elif 'Podpis:' in stripped:
                pass
            elif stripped == '' or stripped.startswith('*Protokół'):
                pass

        i += 1

    return (sigs if sigs else None), i


def convert_md_to_docx(md_path: Path, docx_path: Path):
    doc = Document()
    style_document(doc)

    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    total = len(lines)
    i = 0

    in_signatures = False

    while i < total:
        line = lines[i].rstrip()

        if not line:
            i += 1
            continue

        if line.startswith('# '):
            add_title(doc, line[2:].strip())
            i += 1
            continue

        if line.startswith('## Strony'):
            heading = line[3:].strip()
            wyk, zam, end_idx = parse_parties_block(lines, i + 1)
            add_h1(doc, heading)
            if wyk or zam:
                add_parties_table(doc, wyk, zam)
                i = end_idx
                continue
            i += 1
            continue

        if line.startswith('## '):
            add_h1(doc, line[3:].strip())
            i += 1
            continue

        if line.startswith('### '):
            add_h2(doc, line[4:].strip())
            i += 1
            continue

        if line.startswith('---'):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            i += 1
            continue

        stripped = line.lstrip()
        leading = len(line) - len(stripped)

        if stripped == PAGEBREAK:
            doc.add_page_break()
            i += 1
            continue

        if SIG_ROLE_RE.match(stripped) and not in_signatures:
            sigs, end_idx = detect_signature_block(lines, i)
            if sigs and len(sigs) >= 1:
                add_signatures_table(doc, sigs)
                i = end_idx
                continue

        if stripped.startswith('|') and i + 1 < total and '---' in lines[i + 1]:
            table_lines = []
            while i < total and lines[i].lstrip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            add_table_from_md(doc, table_lines)
            continue

        if stripped.startswith('- [ ]'):
            text = stripped[5:].strip()
            add_checkbox(doc, text)
            i += 1
            continue

        if stripped.startswith('- '):
            text = stripped[2:].strip()
            indent_level = 0 if leading == 0 else 1
            add_bullet(doc, text, indent_level=indent_level)
            i += 1
            continue

        m = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if m:
            text = m.group(2)
            if leading == 0:
                indent_level = 0
                text = f'{m.group(1)}.  {text}'
            else:
                indent_level = 1
                num = int(m.group(1))
                letter = chr(ord('a') + (num - 1) % 26)
                text = f'{letter})  {text}'
            add_numbered(doc, text, indent_level=indent_level)
            i += 1
            continue

        add_paragraph(doc, line)
        i += 1

    doc.save(str(docx_path))
    print(f'✓ {docx_path.name}')


if __name__ == '__main__':
    base = Path(__file__).parent
    files = [
        ('Umowa-wykonanie-strony-Willa-Wojcik.md', 'Umowa-wykonanie-strony-Willa-Wojcik.docx'),
        ('Umowa-powierzenia-Willa-Wojcik.md', 'Umowa-powierzenia-Willa-Wojcik.docx'),
    ]
    for src, dst in files:
        convert_md_to_docx(base / src, base / dst)
